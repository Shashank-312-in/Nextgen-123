import io
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from sms_app.services import attendance_service as ats
from sms_app.services import timetable_service as tts
from api import routes_timetable as rt


class FakeResult:
    def __init__(self, rows=None, rowcount=1, lastrowid=77):
        self.rows = list(rows or [])
        self.rowcount = rowcount
        self.lastrowid = lastrowid

    def fetchone(self):
        return self.rows[0] if self.rows else None

    def fetchall(self):
        return self.rows


class FakeConnection:
    def __init__(self, *, semester=True, student_rows=None, subjects=None):
        self.semester = semester
        self.student_rows = student_rows or {}
        self.subjects = subjects or []
        self.last_sql = ""
        self.last_params = ()
        self.audit_calls = []

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def execute(self, sql, params=()):
        self.last_sql = " ".join(str(sql).split()).lower()
        self.last_params = params
        if self.last_sql.startswith("select id,code,name from academic_semesters"):
            return FakeResult([{"id": 3, "code": "II-I", "name": "II B.Tech I Semester"}] if self.semester else [])
        if self.last_sql.startswith("select id, code, name from subjects"):
            return FakeResult(self.subjects)
        if self.last_sql.startswith("select roll_no, name from students"):
            return FakeResult([self.student_rows.get(params[0])] if self.student_rows.get(params[0]) else [])
        if self.last_sql.startswith("select id from attendance_sessions"):
            return FakeResult([])
        if self.last_sql.startswith("select code, name from subjects where id="):
            return FakeResult([next((s for s in self.subjects if s["id"] == params[0]), None)] if next((s for s in self.subjects if s["id"] == params[0]), None) else [])
        if self.last_sql.startswith("insert into attendance_records"):
            return FakeResult([], rowcount=1)
        if self.last_sql.startswith("update attendance_sessions"):
            return FakeResult([], rowcount=1)
        return FakeResult([], rowcount=1)


def workbook_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def test_status_tolerance_and_invalid_values():
    assert ats._parse_import_status("P", row_no=2) == "Present"
    assert ats._parse_import_status("0", row_no=2) == "Absent"
    assert ats._parse_import_status("yes", row_no=2) == "Present"
    assert ats._parse_import_status("NO", row_no=2) == "Absent"
    with pytest.raises(ValueError, match="status must"):
        ats._parse_import_status("Late", row_no=2)


def test_date_future_check_uses_application_local_date(monkeypatch):
    class FrozenNow:
        @staticmethod
        def date():
            return date(2026, 9, 19)

    monkeypatch.setattr(ats, "local_now", lambda: FrozenNow())
    assert ats._parse_import_date("19-09-2026", row_no=2) == "2026-09-19"
    with pytest.raises(ValueError, match="cannot be in the future"):
        ats._parse_import_date("20-09-2026", row_no=2)



def test_excel_serial_date_is_accepted(monkeypatch):
    class FrozenNow:
        @staticmethod
        def date():
            return date(2026, 9, 19)

    monkeypatch.setattr(ats, "local_now", lambda: FrozenNow())
    # Excel serial 46284 corresponds to 2026-09-19 in the 1900 date system.
    assert ats._parse_import_date(46284, row_no=2) == "2026-09-19"


def test_hod_scope_student_lookup_is_fail_closed():
    class ScopeProbe:
        def __init__(self):
            self.sql = ""

        def execute(self, sql, params=()):
            self.sql = str(sql)
            return FakeResult([])

    probe = ScopeProbe()
    assert ats._hod_scoped_student(probe, roll_no="S1", hod_username="hod1") is None
    normalized = " ".join(probe.sql.split()).lower()
    assert "active=1" in normalized
    assert "coalesce(hod_username,'')" in normalized
    assert "or hod_username is null" not in normalized

def test_lab_duration_is_forced_to_three_hours():
    assert ats._parse_import_duration("1", "LAB", row_no=2) == 3
    assert ats._parse_import_duration(None, "CLASS", row_no=2) == 1
    assert ats._parse_import_duration("3", "CLASS", row_no=2) == 3
    with pytest.raises(ValueError, match="duration"):
        ats._parse_import_duration("4", "CLASS", row_no=2)


def test_subject_name_ambiguity_fails_closed():
    fake = FakeConnection(subjects=[
        {"id": 1, "code": "CS01", "name": "Programming"},
        {"id": 2, "code": "CS02", "name": "Programming"},
    ])
    with pytest.raises(ValueError, match="ambiguous"):
        ats._resolve_import_subject(fake, semester_id=3, subject_code="", subject_name="Programming", row_no=2)


def test_all_unregistered_rows_are_reported_without_partial_session_creation(monkeypatch):
    fake = FakeConnection(
        semester=True,
        student_rows={},
        subjects=[{"id": 1, "code": "CS01", "name": "Programming"}],
    )
    monkeypatch.setattr(ats, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ats, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["Roll No", "Date", "Subject Code", "Status"],
        ["UNKNOWN-1", "2026-09-18", "CS01", "P"],
    ])
    result = ats.upload_attendance_excel(raw=raw, filename="attendance.xlsx", semester_id=3, hod_username="hod1")
    assert result["sessions_created"] == 0
    assert result["records_written"] == 0
    assert result["skipped_count"] == 1
    assert result["skipped_students"][0]["roll_no"] == "UNKNOWN-1"



def test_override_api_guard_is_hod_only():
    with pytest.raises(Exception, match="HOD access required"):
        rt._require_hod(SimpleNamespace(role="ADMIN"))


def test_override_service_rejects_non_hod_before_db(monkeypatch):
    def fail_connect(*args, **kwargs):
        raise AssertionError("non-HOD override call must fail before DB access")

    monkeypatch.setattr(tts, "connect", fail_connect)
    with pytest.raises(PermissionError, match="HOD approval is required"):
        tts.create_override(
            timetable_entry_id=1,
            override_date=date(2026, 9, 21),
            substitute_faculty_username="faculty2",
            reason="Substitute",
            approved_by="admin",
            actor_role="ADMIN",
        )


def test_override_delete_service_rejects_non_hod_before_db(monkeypatch):
    def fail_connect(*args, **kwargs):
        raise AssertionError("non-HOD override delete must fail before DB access")

    monkeypatch.setattr(tts, "connect", fail_connect)
    with pytest.raises(PermissionError, match="HOD access required"):
        tts.delete_override(override_id=1, actor_username="admin", actor_role="ADMIN")


def test_notification_claim_is_same_day_and_not_stale(monkeypatch):
    from zoneinfo import ZoneInfo
    fixed = datetime(2026, 9, 19, 9, 2, 0, tzinfo=ZoneInfo("Asia/Kolkata"))

    class FakeConn:
        def __init__(self):
            self.sql_calls = []

        def __enter__(self):
            return self

        def __exit__(self, *_):
            return False

        def execute(self, sql, params=()):
            normalized = " ".join(str(sql).split())
            self.sql_calls.append((normalized, params))
            if normalized.startswith("SELECT n.*, t.section_name"):
                return FakeResult([{"id": 1, "faculty_username": "faculty1", "timetable_entry_id": 9, "occurrence_date": "2026-09-19", "scheduled_for": "2026-09-19 09:00:00", "section_name": "A", "academic_year": "2026-27", "subject_code": "CS01", "subject_name": "Programming", "custom_label": "", "room": "2"}])
            return FakeResult([], rowcount=1)

    fake = FakeConn()
    monkeypatch.setattr(tts, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(tts, "local_now", lambda: fixed)
    claimed = tts.claim_due_notifications(faculty_username="faculty1", limit=1)
    assert len(claimed) == 1
    select_sql, params = fake.sql_calls[0]
    assert "n.occurrence_date=%s" in select_sql
    assert "n.scheduled_for between %s and %s" in select_sql.lower()
    assert params[1] == fixed.date()
    assert params[2] == fixed.replace(tzinfo=None) - timedelta(minutes=3)
    assert params[3] == fixed.replace(tzinfo=None)

def test_period_map_is_section_relative_and_preserves_afternoon_slots():
    periods = [
        {"key": "m1", "label": "P1", "start": "09:00", "end": "09:50", "section": "MORNING"},
        {"key": "m2", "label": "P2", "start": "09:50", "end": "10:40", "section": "MORNING"},
        {"key": "a1", "label": "P5", "start": "13:20", "end": "14:10", "section": "AFTERNOON"},
        {"key": "a2", "label": "P6", "start": "14:10", "end": "15:00", "section": "AFTERNOON"},
    ]
    mapped = tts._period_map(__import__("json").dumps(periods))
    assert mapped[("MORNING", 0)]["start"] == "09:00"
    assert mapped[("AFTERNOON", 0)]["start"] == "13:20"
    assert mapped[("AFTERNOON", 1)]["label"] == "P6"


def test_duplicate_student_rows_fail_before_any_session_write(monkeypatch):
    fake = FakeConnection(
        semester=True,
        student_rows={"S1": {"roll_no": "S1", "name": "Student One"}},
        subjects=[{"id": 1, "code": "CS01", "name": "Programming"}],
    )
    monkeypatch.setattr(ats, "connect", lambda *args, **kwargs: fake)
    monkeypatch.setattr(ats, "audit", lambda *args, **kwargs: None)
    raw = workbook_bytes([
        ["Roll No", "Date", "Subject Code", "Status"],
        ["S1", "2026-09-18", "CS01", "P"],
        ["S1", "2026-09-18", "CS01", "A"],
    ])
    with pytest.raises(ValueError, match="duplicate attendance row"):
        ats.upload_attendance_excel(raw=raw, filename="attendance.xlsx", semester_id=3, hod_username="hod1")


def test_template_contains_import_contract():
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(ats.build_attendance_template()), data_only=True)
    attendance = workbook["Attendance"]
    guide = workbook["Instructions"]
    assert [attendance.cell(1, c).value for c in range(1, 9)] == [
        "Roll No", "Date", "Subject Code", "Subject Name", "Status", "Session Type", "Duration Hours", "Topic"
    ]
    assert attendance.cell(2, 5).value == "P"
    assert any(guide.cell(r, 1).value == "Status" for r in range(2, guide.max_row + 1))
    workbook.close()


def test_effective_schedule_uses_end_of_spanning_block(monkeypatch):
    class Context:
        def __enter__(self):
            return self
        def __exit__(self, *_):
            return False
        def execute(self, sql, params=()):
            return FakeResult([{
                "timetable_entry_id": 9,
                "timetable_id": 4,
                "day_of_week": "SAT",
                "section": "MORNING",
                "start_slot": 0,
                "duration": 2,
                "block_type": "THEORY",
                "subject_id": 6,
                "custom_label": "",
                "regular_faculty_username": "faculty1",
                "room": "Room 2",
                "semester_id": 3,
                "section_name": "A",
                "academic_year": "2026-27",
                "hod_username": "hod1",
                "period_config_json": __import__("json").dumps([
                    {"key": "m1", "label": "P1", "start": "09:00", "end": "09:50", "section": "MORNING"},
                    {"key": "m2", "label": "P2", "start": "09:50", "end": "10:40", "section": "MORNING"},
                ]),
                "subject_code": "CS01",
                "subject_name": "Programming",
                "override_id": None,
                "substitute_faculty_username": None,
                "reason": None,
                "approved_by": None,
                "override_created_at": None,
                "substitute_faculty_name": None,
                "regular_faculty_name": "Faculty One",
            }])
    monkeypatch.setattr(tts, "connect", lambda *args, **kwargs: Context())
    rows = tts.resolve_effective_schedule(faculty_username="faculty1", target_date=date(2026, 9, 19))
    assert len(rows) == 1
    assert rows[0]["start_time"] == "09:00"
    assert rows[0]["end_time"] == "10:40"
